import sys
import requests
from PyQt5.QtWidgets import (
    QApplication, QMainWindow, QStackedWidget, QVBoxLayout, QLabel,
    QPushButton, QListWidget, QListWidgetItem, QWidget, QMessageBox, QHBoxLayout, 
    QComboBox, QTableWidget, QTableWidgetItem, QLineEdit, QHeaderView, QCheckBox, 
    QTextEdit, QFileDialog, QTabWidget, QSpinBox, QGroupBox, QColorDialog, QDialog
)
from PyQt5.QtCore import Qt
import pandas as pd
from openpyxl import load_workbook
import xlrd
import os

BACKEND_URL = "http://127.0.0.1:5000"

class InitialPage(QWidget):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.layout = QVBoxLayout()
        
        # Título
        title_label = QLabel("Selecciona el tipo de reporte a generar:")
        title_label.setStyleSheet("font-size: 14pt; font-weight: bold;")
        self.layout.addWidget(title_label)
        
        # Botones para seleccionar tipo de reporte
        self.latex_button = QPushButton("Reporte LaTeX")
        self.word_button = QPushButton("Reporte Word")
        
        # Estilizar botones
        button_style = """
            QPushButton {
                padding: 15px;
                font-size: 12pt;
                min-width: 200px;
                margin: 10px;
            }
        """
        self.latex_button.setStyleSheet(button_style)
        self.word_button.setStyleSheet(button_style)
        
        # Agregar botones al layout
        self.layout.addWidget(self.latex_button)
        self.layout.addWidget(self.word_button)
        
        # Conectar eventos
        self.latex_button.clicked.connect(self.select_latex)
        self.word_button.clicked.connect(self.select_word)
        
        self.setLayout(self.layout)

    def select_latex(self):
        # Obtener referencia a MainWindow
        main_window = self.parent().parent()
        main_window.report_type = "latex"
        # Ir a la página de selección de plantillas
        self.parent().setCurrentIndex(1)

    def select_word(self):
        # Obtener referencia a MainWindow
        main_window = self.parent().parent()
        main_window.report_type = "word"
        # Ir directamente a la página de selección de bases de datos
        self.parent().setCurrentIndex(2)

class FirstPage(QWidget):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.layout = QVBoxLayout()
        
        # Botón para seleccionar la carpeta de plantillas
        self.select_folder_button = QPushButton("Seleccionar Carpeta de Plantillas")
        self.select_folder_button.clicked.connect(self.select_template_folder)
        
        self.list_widget = QListWidget()
        self.next_button = QPushButton("Siguiente")
        self.status_label = QLabel("Selecciona una carpeta de plantillas para comenzar.")
        
        # Añadir diccionario para almacenar las plantillas
        self.templates = {}  # Para almacenar {nombre: ruta}

        self.layout.addWidget(self.select_folder_button)
        self.layout.addWidget(QLabel("Selecciona una Plantilla:"))
        self.layout.addWidget(self.list_widget)
        self.layout.addWidget(self.next_button)
        self.layout.addWidget(self.status_label)
        self.setLayout(self.layout)

        self.next_button.clicked.connect(self.go_to_next_page)

    def select_template_folder(self):
        """Permite al usuario seleccionar la carpeta de plantillas."""
        folder_path = QFileDialog.getExistingDirectory(self, "Seleccionar Carpeta de Plantillas", os.path.expanduser("~"), QFileDialog.ShowDirsOnly)
        
        if folder_path:
            # Verificar si la carpeta contiene la subcarpeta "Plantilla para reportes"
            default_folder = os.path.join(folder_path, "Plantilla para reportes")
            if os.path.exists(default_folder):
                folder_path = default_folder
            
            # Enviar la ruta de la carpeta al backend
            response = requests.post(f"{BACKEND_URL}/set-template-folder", json={'folder_path': folder_path})
            
            if response.status_code == 200:
                self.status_label.setText(f"Carpeta seleccionada: {folder_path}")
                self.load_templates()
            else:
                self.status_label.setText("Error al seleccionar la carpeta de plantillas.")
                QMessageBox.warning(self, "Error", "No se pudo configurar la carpeta de plantillas.")

    def load_templates(self):
        """Carga las plantillas desde la carpeta seleccionada."""
        try:
            response = requests.get(f"{BACKEND_URL}/templates")
            if response.status_code == 200:
                templates = response.json()
                self.templates = {template["name"]: template["path"] for template in templates}
                self.list_widget.clear()
                for template_name in self.templates:
                    self.list_widget.addItem(template_name)
                self.status_label.setText("Plantillas cargadas correctamente.")
            else:
                self.status_label.setText("Error al cargar plantillas.")
        except Exception as e:
            self.status_label.setText(f"Error: {e}")

    def go_to_next_page(self):
        selected_item = self.list_widget.currentItem()
        if selected_item:
            template_name = selected_item.text()
            template_path = self.templates[template_name]
            
            # Enviar la plantilla seleccionada al backend
            response = requests.post(f"{BACKEND_URL}/set-template", json={'template_path': template_path})
            
            if response.status_code == 200:
                # Obtener referencia a MainWindow y guardar la ruta completa
                main_window = self.parent().parent()
                main_window.selected_template = template_path
                
                print(f"Plantilla seleccionada: {template_path}")  # Para depuración
                # Cambiar a la página de bases de datos (índice 2)
                self.parent().setCurrentIndex(2)
            else:
                QMessageBox.warning(self, "Error", "No se pudo configurar la plantilla seleccionada.")
        else:
            QMessageBox.warning(self, "Advertencia", "Por favor, selecciona una plantilla.")

    def go_to_previous_page(self):
        # Corregir el índice para volver a la página anterior (índice 0)
        self.parent().setCurrentIndex(0)
            
class SecondPage(QWidget):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.layout = QVBoxLayout()
        
        # Botón para seleccionar la carpeta de bases de datos
        self.select_folder_button = QPushButton("Seleccionar Carpeta de Bases de Datos")
        self.select_folder_button.clicked.connect(self.select_database_folder)
        
        # Barra de búsqueda
        self.search_layout = QHBoxLayout()
        self.search_label = QLabel("Buscar:")
        self.search_input = QLineEdit()
        self.search_input.setPlaceholderText("Escriba para filtrar bases de datos...")
        self.search_input.textChanged.connect(self.filter_databases)
        self.search_layout.addWidget(self.search_label)
        self.search_layout.addWidget(self.search_input)
        
        self.list_widget = QListWidget()
        self.back_button = QPushButton("Regresar")
        self.next_button = QPushButton("Seleccionar Bases de Datos")
        self.status_label = QLabel("Selecciona una carpeta de bases de datos para comenzar.")

        # Añadir widgets al layout principal
        self.layout.addWidget(self.select_folder_button)
        self.layout.addLayout(self.search_layout)
        self.layout.addWidget(QLabel("Selecciona las Bases de Datos:"))
        self.layout.addWidget(self.list_widget)
        self.layout.addWidget(self.next_button)
        self.layout.addWidget(self.back_button)
        self.layout.addWidget(self.status_label)
        self.setLayout(self.layout)

        # Botones
        self.back_button.clicked.connect(self.go_to_previous_page)

        # Lista para almacenar todas las bases de datos
        self.all_databases = []
        # Diccionario para mantener el estado de las casillas
        self.checked_states = {}
        
        # Conectar el evento de cambio de estado de las casillas
        self.list_widget.itemChanged.connect(self.on_item_changed)

    def select_database_folder(self):
        """Permite al usuario seleccionar la carpeta de bases de datos."""
        folder_path = QFileDialog.getExistingDirectory(self, "Seleccionar Carpeta de Bases de Datos", os.path.expanduser("~"), QFileDialog.ShowDirsOnly)
        
        if folder_path:
            # Enviar la ruta de la carpeta al backend
            response = requests.post(f"{BACKEND_URL}/set-database-folder", json={'folder_path': folder_path})
            
            if response.status_code == 200:
                self.status_label.setText(f"Carpeta seleccionada: {folder_path}")
                self.load_databases()
            else:
                self.status_label.setText("Error al seleccionar la carpeta de bases de datos.")
                QMessageBox.warning(self, "Error", "No se pudo configurar la carpeta de bases de datos.")

    def load_databases(self):
        """Carga las bases de datos desde la carpeta seleccionada."""
        try:
            response = requests.get(f"{BACKEND_URL}/databases")
            if response.status_code == 200:
                databases = response.json()
                self.all_databases = databases
                self.list_widget.clear()
                self.checked_states = {}
                for db in databases:
                    full_text = f"{db['name']} ({db['path']})"
                    item = QListWidgetItem(full_text)
                    item.setCheckState(Qt.Unchecked)
                    self.checked_states[full_text] = Qt.Unchecked
                    self.list_widget.addItem(item)
                self.status_label.setText("Bases de datos cargadas correctamente.")
            else:
                self.status_label.setText("Error al cargar bases de datos.")
        except Exception as e:
            self.status_label.setText(f"Error: {e}")

    def on_item_changed(self, item):
        """Actualiza el estado global de las casillas cuando cambia una selección."""
        self.checked_states[item.text()] = item.checkState()

    def filter_databases(self, text):
        """Filtra las bases de datos según el texto de búsqueda y mantiene las selecciones."""
        self.list_widget.clear()
        search_text = text.lower()
        
        for db in self.all_databases:
            full_text = f"{db['name']} ({db['path']})"
            if search_text in db['name'].lower():
                item = QListWidgetItem(full_text)
                # Usar el estado guardado o establecer como desmarcado si no existe
                item.setCheckState(self.checked_states.get(full_text, Qt.Unchecked))
                self.list_widget.addItem(item)

    def store_selected_databases(self):
        """Almacena todas las bases de datos seleccionadas usando el estado global."""
        selected_databases = []
        
        # Usar checked_states para obtener todas las selecciones
        for db in self.all_databases:
            full_text = f"{db['name']} ({db['path']})"
            if self.checked_states.get(full_text) == Qt.Checked:
                selected_databases.append(full_text)
                print(f"Base seleccionada: {full_text}")  # Depuración

        if selected_databases:
            # Obtener la referencia a MainWindow
            main_window = self.parent().parent()
            main_window.selected_databases = selected_databases.copy()  # Hacer una copia de la lista
            print(f"Bases guardadas en MainWindow: {main_window.selected_databases}")  # Depuración
            
            QMessageBox.information(
                self, 
                "Bases Seleccionadas", 
                f"Bases seleccionadas:\n{', '.join(selected_databases)}"
            )
            # Llamar a start_iteration desde MainWindow
            main_window.start_iteration()
            # Cambiar a la página de selección de hoja (índice 3)
            self.parent().setCurrentIndex(3)
        else:
            QMessageBox.warning(self, "Advertencia", "Por favor, selecciona al menos una base de datos.")

    def go_to_previous_page(self):
        self.parent().setCurrentIndex(1)  # Volver a la página de plantillas

class ThirdPage(QWidget):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.layout = QVBoxLayout()
        
        # Componentes de la interfaz
        self.file_label = QLabel("Archivo seleccionado:")
        self.sheet_selector = QComboBox()
        self.table_preview = QTableWidget()
        self.header_input = QLineEdit()
        self.preview_button = QPushButton("Previsualizar Hoja")
        self.next_button = QPushButton("Seleccionar Encabezado")
        self.back_button = QPushButton("Regresar")

        # Configurar diseño
        self.layout.addWidget(self.file_label)
        self.layout.addWidget(QLabel("Selecciona una hoja del archivo:"))
        self.layout.addWidget(self.sheet_selector)
        self.layout.addWidget(self.preview_button)
        self.layout.addWidget(QLabel("Vista previa de la hoja:"))
        self.layout.addWidget(self.table_preview)
        self.layout.addWidget(QLabel("Selecciona la fila de encabezados (base 1):"))
        self.layout.addWidget(self.header_input)
        self.layout.addWidget(self.next_button)
        self.layout.addWidget(self.back_button)
        self.setLayout(self.layout)

        # Eventos de botones
        self.preview_button.clicked.connect(self.preview_sheet)
        self.next_button.clicked.connect(self.select_header_row)
        self.back_button.clicked.connect(self.go_to_previous_page)

        # Variables
        self.selected_file = None
        self.selected_sheet = None
        self.selected_header_row = None

    def load_file(self, file_path):
        """Carga las hojas del archivo Excel seleccionado en el ComboBox."""
        self.selected_file = file_path
        self.file_label.setText(f"Archivo seleccionado: {file_path}")
        self.sheet_selector.clear()

        try:
            extension = os.path.splitext(file_path)[1].lower()
            if extension == ".xlsx":
                workbook = load_workbook(filename=file_path, read_only=True)
                sheets = workbook.sheetnames
            elif extension == ".xls":
                workbook = xlrd.open_workbook(file_path)
                sheets = workbook.sheet_names()
            else:
                QMessageBox.warning(self, "Error", "Formato de archivo no soportado.")
                return

            self.sheet_selector.addItems(sheets)
        except Exception as e:
            QMessageBox.warning(self, "Error", f"No se pudo cargar el archivo: {e}")

    def preview_sheet(self):
        """Carga y muestra una vista previa de la hoja seleccionada."""
        self.selected_sheet = self.sheet_selector.currentText()
        if not self.selected_file or not self.selected_sheet:
            QMessageBox.warning(self, "Error", "Por favor selecciona un archivo y una hoja.")
            return

        try:
            # Asegurar que header_row tenga un valor por defecto
            header_row = getattr(self, 'selected_header_row', 0)
            if header_row is None:
                header_row = 0
                
            print(f"Enviando solicitud con:")
            print(f"- Archivo: {self.selected_file}")
            print(f"- Hoja: {self.selected_sheet}")
            print(f"- Header row: {header_row}")
            
            response = requests.post(f"{BACKEND_URL}/preview-sheet", json={
                'file_path': self.selected_file,
                'sheet_name': str(self.selected_sheet),
                'header_row': header_row  # Ahora siempre enviará un número
            })
            
            if response.status_code == 200:
                data = response.json()
                self.show_table_from_data(data['columns'], data['data'])
            else:
                error_msg = response.json().get('error', 'Error desconocido')
                QMessageBox.warning(self, "Error", f"No se pudo cargar la hoja. Error: {error_msg}")
                print(f"Error response: {response.text}")
        except Exception as e:
            import traceback
            print(traceback.format_exc())
            QMessageBox.warning(self, "Error", f"No se pudo cargar la hoja: {str(e)}")

    def show_table(self, df):
        """Muestra un DataFrame en la tabla de vista previa."""
        self.table_preview.setColumnCount(len(df.columns))
        self.table_preview.setRowCount(len(df))

        # Agregar encabezados
        self.table_preview.setHorizontalHeaderLabels(df.columns)

        # Agregar datos
        for i in range(len(df)):
            for j in range(len(df.columns)):
                self.table_preview.setItem(i, j, QTableWidgetItem(str(df.iloc[i, j])))

        # Ajustar el ancho de las columnas
        self.table_preview.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)

    def select_header_row(self):
        """Selecciona la fila de encabezados y navega a la siguiente página."""
        try:
            header_row = int(self.header_input.text()) - 1
            if header_row < 0:
                QMessageBox.warning(self, "Error", "La fila debe ser mayor a 0.")
                return

            self.selected_header_row = header_row
            
            # Leer los nombres de las columnas con el encabezado seleccionado
            df = pd.read_excel(
                self.selected_file,
                sheet_name=self.selected_sheet,
                header=header_row,
                nrows=1  # Solo leer la fila de encabezados
            )
            
            # Almacenar los nombres de las columnas (convertidos a string)
            self.column_names = [str(col) for col in df.columns.tolist()]
            
            # Configurar la cuarta página con los nombres de las columnas
            fourth_page = self.parent().widget(4)  # Corregido el índice
            fourth_page.load_sheet(
                self.selected_file,
                self.selected_sheet,
                self.column_names
            )
            
            QMessageBox.information(
                self, 
                "Fila Seleccionada", 
                f"Fila de encabezados seleccionada: {header_row + 1}\n"
                f"Columnas detectadas: {', '.join(self.column_names)}"
            )
            
            # Navegar a la cuarta página
            self.parent().setCurrentIndex(4)  # Corregido el índice
            
        except ValueError:
            QMessageBox.warning(self, "Error", "Por favor ingresa un número válido.")

    def go_to_previous_page(self):
        self.parent().setCurrentIndex(2)

    def show_table_from_data(self, columns, data):
        """Muestra los datos recibidos del backend en la tabla."""
        self.table_preview.setColumnCount(len(columns))
        self.table_preview.setRowCount(len(data))
        
        # Establecer encabezados
        self.table_preview.setHorizontalHeaderLabels(columns)
        
        # Llenar datos
        for row_idx, row_data in enumerate(data):
            for col_idx, column in enumerate(columns):
                item = QTableWidgetItem(str(row_data[column]))
                self.table_preview.setItem(row_idx, col_idx, item)
        
        # Ajustar el ancho de las columnas
        self.table_preview.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)

    def clear_selections(self):
        """Limpia las selecciones anteriores."""
        self.selected_file = None
        self.selected_sheet = None
        self.selected_header_row = None
        self.sheet_selector.clear()
        self.header_input.clear()
        self.table_preview.setRowCount(0)
        self.table_preview.setColumnCount(0)

class FourthPage(QWidget):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.layout = QVBoxLayout()

        # Componentes de la interfaz
        self.sheet_label = QLabel("Hoja seleccionada:")
        self.table_preview = QTableWidget()
        self.graph_selector_label = QLabel("Selecciona los gráficos a realizar:")
        self.bar_chart_checkbox = QCheckBox("Bar Chart")
        self.time_series_checkbox = QCheckBox("Time Series")
        self.column_selector_label = QLabel("Selecciona las columnas para el gráfico:")
        self.x_column_dropdown = QComboBox()
        self.y_column_dropdown = QComboBox()
        self.save_button = QPushButton("Guardar Selección")
        self.back_button = QPushButton("Regresar")
        self.edit_graph_button = QPushButton("Editar Gráfico")

        # Configuración del layout
        self.layout.addWidget(self.sheet_label)
        self.layout.addWidget(self.table_preview)
        self.layout.addWidget(self.graph_selector_label)
        self.layout.addWidget(self.bar_chart_checkbox)
        self.layout.addWidget(self.time_series_checkbox)
        self.layout.addWidget(self.column_selector_label)
        self.layout.addWidget(QLabel("Columna para el eje X:"))
        self.layout.addWidget(self.x_column_dropdown)
        self.layout.addWidget(QLabel("Columna para el eje Y:"))
        self.layout.addWidget(self.y_column_dropdown)
        self.layout.addWidget(self.save_button)
        self.layout.addWidget(self.edit_graph_button)
        self.layout.addWidget(self.back_button)
        self.setLayout(self.layout)

        # Eventos
        self.save_button.clicked.connect(self.save_selections)
        self.edit_graph_button.clicked.connect(self.edit_graph)
        self.back_button.clicked.connect(self.go_to_previous_page)

        # Variables
        self.selected_sheet = None
        self.selected_file = None
        self.selected_graphs = []
        self.selected_columns = {}
        self.graph_options = None

    def load_sheet(self, file_path, sheet_name, column_names):
        """Carga la información de la hoja y configura los dropdowns de columnas."""
        self.selected_file = file_path
        self.selected_sheet = sheet_name
        self.sheet_label.setText(f"Hoja seleccionada: {sheet_name}")
        
        try:
            # Leer los datos para mostrar en la tabla
            df = pd.read_excel(
                file_path,
                sheet_name=sheet_name,
                header=self.parent().parent().third_page.selected_header_row
            )
            
            # Mostrar los datos en la tabla
            self.table_preview.setColumnCount(len(column_names))
            self.table_preview.setRowCount(len(df))
            
            # Establecer encabezados
            self.table_preview.setHorizontalHeaderLabels(column_names)
            
            # Llenar datos usando iloc para acceder por posición
            for i in range(len(df)):
                for j in range(len(df.columns)):
                    item = QTableWidgetItem(str(df.iloc[i, j]))  # Usar iloc para acceder por posición
                    self.table_preview.setItem(i, j, item)
            
            # Ajustar el ancho de las columnas
            self.table_preview.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
            
            # Actualizar los dropdowns con los nombres de las columnas
            self.x_column_dropdown.clear()
            self.y_column_dropdown.clear()
            self.x_column_dropdown.addItems(column_names)
            self.y_column_dropdown.addItems(column_names)
            
        except Exception as e:
            print(f"Error al cargar la hoja en la cuarta página: {str(e)}")
            import traceback
            print(traceback.format_exc())
            QMessageBox.warning(self, "Error", f"No se pudo cargar la vista previa: {str(e)}")

    def show_table(self, df):
        """Muestra un DataFrame en la tabla de vista previa."""
        self.table_preview.setColumnCount(len(df.columns))
        self.table_preview.setRowCount(len(df))

        # Agregar encabezados
        self.table_preview.setHorizontalHeaderLabels(df.columns)

        # Agregar datos
        for i in range(len(df)):
            for j in range(len(df.columns)):
                self.table_preview.setItem(i, j, QTableWidgetItem(str(df.iloc[i, j])))

        # Ajustar el ancho de las columnas
        self.table_preview.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)

    def edit_graph(self):
        dialog = GraphEditorDialog(self)
        if dialog.exec_() == QDialog.Accepted:
            self.graph_options = dialog.get_options()
            QMessageBox.information(self, "Opciones Guardadas", 
                                  "Las opciones de personalización han sido guardadas.")

    def save_selections(self):
        """Guarda las selecciones realizadas para los gráficos y columnas."""
        # Validar que al menos un gráfico esté seleccionado
        if not (self.bar_chart_checkbox.isChecked() or self.time_series_checkbox.isChecked()):
            QMessageBox.warning(self, "Advertencia", "Por favor selecciona al menos un tipo de gráfico.")
            return

        # Obtener gráficos seleccionados
        self.selected_graphs = []
        if self.bar_chart_checkbox.isChecked():
            self.selected_graphs.append("bar_chart")
        if self.time_series_checkbox.isChecked():
            self.selected_graphs.append("time_series")

        # Validar selección de columnas
        x_column = self.x_column_dropdown.currentText()
        y_column = self.y_column_dropdown.currentText()
        if not x_column or not y_column:
            QMessageBox.warning(self, "Advertencia", "Por favor selecciona las columnas para los gráficos.")
            return

        # Guardar selección de columnas
        self.selected_columns = {
            "x_column": x_column, 
            "y_column": y_column,
            "graph_options": self.graph_options
        }

        # Obtener referencia a MainWindow y guardar las decisiones
        main_window = self.parent().parent()
        main_window.save_graph_decisions()

    def go_to_previous_page(self):
        self.parent().setCurrentIndex(3)

    def clear_selections(self):
        """Limpia las selecciones de gráficos y columnas pero mantiene la hoja seleccionada."""
        self.x_column_dropdown.clear()
        self.y_column_dropdown.clear()
        self.bar_chart_checkbox.setChecked(False)
        self.time_series_checkbox.setChecked(False)
        self.selected_graphs = []
        self.selected_columns = {}
        self.graph_options = None
        # No limpiamos self.selected_sheet ni self.selected_file

class FinalPage(QWidget):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.layout = QVBoxLayout()

        # Componentes de la interfaz
        self.summary_label = QLabel("Resumen de Decisiones:")
        self.summary_text = QTextEdit()
        self.summary_text.setReadOnly(True)
        self.generate_button = QPushButton("Generar Reporte")
        self.back_button = QPushButton("Regresar")
        self.clear_button = QPushButton("Limpiar Selecciones")
        self.output_label = QLabel("Directorio de salida:")
        
        # Directorio de salida por defecto
        self.output_directory = os.path.join(os.path.expanduser("~"), "Desktop", "Reporte Generado")

        # Configuración del layout
        self.layout.addWidget(self.summary_label)
        self.layout.addWidget(self.summary_text)
        self.layout.addWidget(self.output_label)
        self.layout.addWidget(self.generate_button)
        self.layout.addWidget(self.back_button)
        self.layout.addWidget(self.clear_button)
        self.setLayout(self.layout)

        # Eventos
        self.generate_button.clicked.connect(self.generate_report)
        self.back_button.clicked.connect(self.go_to_previous_page)
        self.clear_button.clicked.connect(self.clear_all_selections)

        # Variables
        self.decisions = []
        self.template_path = None

    def load_decisions(self, template_path, decisions):
        """Carga las decisiones y la plantilla seleccionada para mostrarlas en la interfaz."""
        self.template_path = template_path
        self.decisions = decisions

        # Mostrar el resumen de las decisiones
        summary = f"Plantilla seleccionada:\n{self.template_path}\n\n"
        for db_decision in self.decisions:
            summary += f"Base de Datos: {db_decision['database']}\n"
            for sheet_decision in db_decision['decisions']:
                summary += f"  Hoja: {sheet_decision['sheet']}\n"
                summary += f"  Fila de encabezados: {sheet_decision['header_row'] + 1}\n"
                summary += f"  Gráficos:\n"
                for graph in sheet_decision['graphs']:
                    summary += f"    - {graph.capitalize()} (X: {sheet_decision['columns']['x_column']}, Y: {sheet_decision['columns']['y_column']})\n"
            summary += "\n"

        self.summary_text.setText(summary)
        self.output_label.setText(f"Directorio de salida: {self.output_directory}")

    def generate_report(self):
        """Genera el reporte basado en las decisiones tomadas."""
        try:
            main_window = self.parent().parent()
            report_type = main_window.report_type

            response = requests.post(f"{BACKEND_URL}/generate-report", json={
                'template_path': self.template_path if report_type == "latex" else None,
                'decisions': self.decisions,
                'report_type': report_type,
                'output_directory': self.output_directory
            })
            
            if response.status_code == 200:
                data = response.json()
                QMessageBox.information(
                    self, 
                    "Reporte Generado", 
                    f"El reporte se ha generado correctamente en:\n{data['output_path']}\n\n" +
                    ("Para el reporte Word, copia el contenido del archivo .txt y las imágenes en un documento de Word." 
                     if report_type == "word" else "")
                )
            else:
                error_msg = response.json().get('error', 'Error desconocido')
                QMessageBox.warning(self, "Error", f"No se pudo generar el reporte: {error_msg}")
        except Exception as e:
            QMessageBox.critical(self, "Error", f"Ocurrió un error al generar el reporte: {str(e)}")

    def go_to_previous_page(self):
        """Regresa a la página anterior y limpia solo las decisiones guardadas."""
        try:
            # Obtener referencia a MainWindow
            main_window = self.parent().parent()
            
            # Limpiar solo las decisiones guardadas
            main_window.database_decisions = []
            main_window.current_database_index = 0
            
            # Limpiar el resumen
            self.summary_text.clear()
            
            # Volver a la página anterior (cuarta página)
            self.parent().setCurrentIndex(4)
            
        except Exception as e:
            print(f"Error al limpiar decisiones: {e}")
            # Aún así intentar volver a la página anterior
            self.parent().setCurrentIndex(4)

    def clear_all_selections(self):
        """Limpia todas las selecciones y reinicia la aplicación."""
        try:
            # Obtener referencia a MainWindow
            main_window = self.parent().parent()
            
            # Limpiar todas las variables de selección
            main_window.report_type = None
            main_window.selected_template = None
            main_window.selected_databases = []
            main_window.database_decisions = []
            main_window.current_database_index = 0
            
            # Limpiar las selecciones en todas las páginas
            main_window.first_page.list_widget.clear()
            main_window.second_page.list_widget.clear()
            main_window.second_page.checked_states = {}
            main_window.third_page.clear_selections()
            main_window.fourth_page.clear_selections()
            
            # Limpiar el resumen
            self.summary_text.clear()
            
            # Mostrar mensaje de confirmación
            QMessageBox.information(
                self,
                "Selecciones Limpiadas",
                "Todas las selecciones han sido limpiadas. La aplicación se reiniciará."
            )
            
            # Volver a la página inicial
            self.parent().setCurrentIndex(0)
            
        except Exception as e:
            print(f"Error al limpiar selecciones: {e}")
            QMessageBox.warning(
                self,
                "Error",
                "Hubo un problema al limpiar las selecciones. Por favor, intente nuevamente."
            )

class GraphEditorDialog(QDialog):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setWindowTitle("Editor de Gráficos")
        self.setModal(True)
        self.setMinimumWidth(600)
        
        layout = QVBoxLayout()
        
        # Crear pestañas
        self.tabs = QTabWidget()
        
        # Pestaña de Colores
        colors_tab = QWidget()
        colors_layout = QVBoxLayout()
        
        self.color_pickers = {}
        for color_type in ['line', 'bar', 'point', 'grid']:
            color_layout = QHBoxLayout()
            color_layout.addWidget(QLabel(f"Color de {color_type}:"))
            color_picker = QPushButton()
            color_picker.setFixedSize(50, 20)
            color_picker.clicked.connect(lambda checked, t=color_type: self.choose_color(t))
            self.color_pickers[color_type] = color_picker
            color_layout.addWidget(color_picker)
            colors_layout.addLayout(color_layout)
        
        colors_tab.setLayout(colors_layout)
        
        # Pestaña de Formato de Datos
        data_tab = QWidget()
        data_layout = QVBoxLayout()
        
        self.data_format = QComboBox()
        self.data_format.addItems(['Número', 'Porcentaje', 'Moneda'])
        data_layout.addWidget(QLabel("Formato de datos:"))
        data_layout.addWidget(self.data_format)
        
        self.decimal_places = QSpinBox()
        self.decimal_places.setRange(0, 10)
        data_layout.addWidget(QLabel("Decimales:"))
        data_layout.addWidget(self.decimal_places)
        
        data_tab.setLayout(data_layout)
        
        # Pestaña de Etiquetas
        labels_tab = QWidget()
        labels_layout = QVBoxLayout()
        
        self.show_values = QCheckBox("Mostrar valores")
        self.show_points = QCheckBox("Mostrar puntos")
        self.label_rotation = QSpinBox()
        self.label_rotation.setRange(0, 90)
        self.label_rotation.setValue(45)
        
        labels_layout.addWidget(self.show_values)
        labels_layout.addWidget(self.show_points)
        labels_layout.addWidget(QLabel("Rotación de etiquetas:"))
        labels_layout.addWidget(self.label_rotation)
        
        labels_tab.setLayout(labels_layout)
        
        # Pestaña de Ejes y Título
        axes_tab = QWidget()
        axes_layout = QVBoxLayout()
        
        self.x_label = QLineEdit()
        self.y_label = QLineEdit()
        self.title = QLineEdit()
        
        axes_layout.addWidget(QLabel("Etiqueta eje X:"))
        axes_layout.addWidget(self.x_label)
        axes_layout.addWidget(QLabel("Etiqueta eje Y:"))
        axes_layout.addWidget(self.y_label)
        axes_layout.addWidget(QLabel("Título:"))
        axes_layout.addWidget(self.title)
        
        # Opciones de fuente
        font_group = QGroupBox("Opciones de fuente")
        font_layout = QVBoxLayout()
        
        self.font_family = QComboBox()
        self.font_family.addItems(['serif', 'sans-serif', 'monospace'])
        self.font_size = QSpinBox()
        self.font_size.setRange(8, 24)
        self.font_size.setValue(12)
        self.font_bold = QCheckBox("Negrita")
        self.font_italic = QCheckBox("Cursiva")
        
        font_layout.addWidget(QLabel("Familia de fuente:"))
        font_layout.addWidget(self.font_family)
        font_layout.addWidget(QLabel("Tamaño:"))
        font_layout.addWidget(self.font_size)
        font_layout.addWidget(self.font_bold)
        font_layout.addWidget(self.font_italic)
        
        font_group.setLayout(font_layout)
        axes_layout.addWidget(font_group)
        
        # Agregar opciones de grid
        grid_group = QGroupBox("Opciones de Grid")
        grid_layout = QVBoxLayout()
        
        self.show_grid = QCheckBox("Mostrar Grid")
        self.grid_style = QComboBox()
        self.grid_style.addItems(['-', '--', ':', '-.'])
        self.grid_alpha = QSpinBox()
        self.grid_alpha.setRange(0, 100)
        self.grid_alpha.setValue(30)
        
        grid_layout.addWidget(self.show_grid)
        grid_layout.addWidget(QLabel("Estilo de Grid:"))
        grid_layout.addWidget(self.grid_style)
        grid_layout.addWidget(QLabel("Transparencia:"))
        grid_layout.addWidget(self.grid_alpha)
        
        grid_group.setLayout(grid_layout)
        axes_layout.addWidget(grid_group)
        
        axes_tab.setLayout(axes_layout)
        
        # Agregar pestañas
        self.tabs.addTab(colors_tab, "Colores")
        self.tabs.addTab(data_tab, "Formato de Datos")
        self.tabs.addTab(labels_tab, "Etiquetas")
        self.tabs.addTab(axes_tab, "Ejes y Título")
        
        layout.addWidget(self.tabs)
        
        # Botones
        buttons_layout = QHBoxLayout()
        self.save_button = QPushButton("Guardar")
        self.cancel_button = QPushButton("Cancelar")
        
        self.save_button.clicked.connect(self.accept)
        self.cancel_button.clicked.connect(self.reject)
        
        buttons_layout.addWidget(self.save_button)
        buttons_layout.addWidget(self.cancel_button)
        
        layout.addLayout(buttons_layout)
        self.setLayout(layout)
        
        # Inicializar valores por defecto
        self.set_default_values()

    def set_default_values(self):
        # Colores por defecto
        default_colors = {
            'line': '#1f77b4',
            'bar': '#1f77b4',
            'point': '#1f77b4',
            'grid': '#cccccc'
        }
        for color_type, color in default_colors.items():
            self.color_pickers[color_type].setStyleSheet(f"background-color: {color}")
        
        # Establecer valores por defecto para las etiquetas
        self.show_values.setChecked(False)  # Desactivar mostrar valores por defecto
        self.show_points.setChecked(False)  # Desactivar mostrar puntos por defecto
        self.label_rotation.setValue(45)    # Rotación por defecto
        
        # Establecer valores por defecto para el grid
        self.show_grid.setChecked(True)     # Activar grid por defecto
        self.grid_style.setCurrentText('--')  # Estilo de grid por defecto
        self.grid_alpha.setValue(30)        # Transparencia por defecto

    def choose_color(self, color_type):
        color = QColorDialog.getColor()
        if color.isValid():
            self.color_pickers[color_type].setStyleSheet(f"background-color: {color.name()}")

    def get_options(self):
        return {
            'colors': {
                'line': self.color_pickers['line'].styleSheet().split(': ')[1],
                'bar': self.color_pickers['bar'].styleSheet().split(': ')[1],
                'point': self.color_pickers['point'].styleSheet().split(': ')[1],
                'grid': self.color_pickers['grid'].styleSheet().split(': ')[1]
            },
            'data_format': {
                'y_axis': self.data_format.currentText().lower(),
                'decimal_places': self.decimal_places.value()
            },
            'labels': {
                'show_values': self.show_values.isChecked(),
                'show_points': self.show_points.isChecked(),
                'rotation': self.label_rotation.value(),
                'font_size': 8  # Tamaño de fuente por defecto
            },
            'axes': {
                'x_label': self.x_label.text(),
                'y_label': self.y_label.text(),
                'title': self.title.text(),
                'title_font': {
                    'family': self.font_family.currentText(),
                    'size': self.font_size.value(),
                    'weight': 'bold' if self.font_bold.isChecked() else 'normal',
                    'style': 'italic' if self.font_italic.isChecked() else 'normal'
                },
                'label_font': {
                    'family': self.font_family.currentText(),
                    'size': self.font_size.value(),
                    'weight': 'normal',
                    'style': 'normal'
                }
            },
            'grid': {
                'show': self.show_grid.isChecked(),
                'style': self.grid_style.currentText(),
                'alpha': self.grid_alpha.value() / 100.0  # Convertir a decimal
            }
        }

class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("Generador de Reportes")
        self.setGeometry(100, 100, 800, 600)

        # Configuración de páginas
        self.stack = QStackedWidget()
        self.initial_page = InitialPage(self.stack)
        self.first_page = FirstPage(self.stack)
        self.second_page = SecondPage(self.stack)
        self.third_page = ThirdPage(self.stack)
        self.fourth_page = FourthPage(self.stack)
        self.final_page = FinalPage(self.stack)

        # Añadir páginas en el orden correcto
        self.stack.addWidget(self.initial_page)  # índice 0
        self.stack.addWidget(self.first_page)    # índice 1
        self.stack.addWidget(self.second_page)   # índice 2
        self.stack.addWidget(self.third_page)    # índice 3
        self.stack.addWidget(self.fourth_page)   # índice 4
        self.stack.addWidget(self.final_page)    # índice 5

        self.setCentralWidget(self.stack)

        # Variables para almacenar selecciones
        self.report_type = None
        self.selected_template = None
        self.selected_databases = []
        self.database_decisions = []
        self.current_database_index = 0

        # Conexión entre páginas
        self.second_page.next_button.clicked.connect(self.second_page.store_selected_databases)

    def save_graph_decisions(self):
        """Guarda las decisiones de gráficos y maneja la iteración de bases de datos."""
        if not hasattr(self, 'database_decisions'):
            self.database_decisions = []

        # Verificar que se hayan hecho las selecciones necesarias
        if not self.fourth_page.selected_graphs or not self.fourth_page.selected_columns:
            QMessageBox.warning(self, "Advertencia", "Por favor, selecciona los gráficos y columnas antes de continuar.")
            return

        # Verificar que tenemos bases de datos seleccionadas
        if not self.selected_databases:
            QMessageBox.warning(self, "Error", "No hay bases de datos seleccionadas.")
            return

        try:
            current_database = self.selected_databases[self.current_database_index]
            database_path = current_database[current_database.find("(")+1:current_database.find(")")]
            
            # Crear o actualizar las decisiones para la base de datos actual
            current_decision = {
                'database': database_path,
                'decisions': [{
                    'sheet': self.third_page.selected_sheet,
                    'header_row': self.third_page.selected_header_row,
                    'graphs': self.fourth_page.selected_graphs,
                    'columns': self.fourth_page.selected_columns
                }]
            }
            
            self.database_decisions.append(current_decision)
            
            # Verificar si hay más bases de datos para procesar
            if self.current_database_index + 1 < len(self.selected_databases):
                # Incrementar el índice solo si hay más bases de datos
                self.current_database_index += 1
                
                # Cargar la siguiente base de datos
                next_database = self.selected_databases[self.current_database_index]
                database_path = next_database[next_database.find("(")+1:next_database.find(")")]
                
                # Limpiar selecciones anteriores
                self.third_page.clear_selections()
                self.fourth_page.clear_selections()
                
                # Cargar la siguiente base de datos en la tercera página
                self.third_page.load_file(database_path)
                
                # Volver a la tercera página
                self.stack.setCurrentIndex(3)
                
                QMessageBox.information(
                    self,
                    "Siguiente Base de Datos",
                    f"Por favor, procese la siguiente base de datos:\n{next_database.split(' (')[0]}"
                )
            else:
                # Si no hay más bases de datos, ir a la página final
                self.go_to_final_page()
                
        except IndexError as e:
            print(f"Error de índice: {e}")
            print(f"Bases de datos seleccionadas: {self.selected_databases}")
            print(f"Índice actual: {self.current_database_index}")
            QMessageBox.warning(self, "Error", "Hubo un problema al procesar las bases de datos. Por favor, intente nuevamente.")
            # Reiniciar el proceso
            self.current_database_index = 0
            self.database_decisions = []
            self.stack.setCurrentIndex(2)  # Volver a la página de selección de bases de datos

    def start_iteration(self):
        """Inicia la iteración de bases de datos seleccionadas."""
        if not self.selected_databases:
            QMessageBox.warning(self, "Advertencia", "Por favor, selecciona al menos una base de datos.")
            return

        print(f"Iniciando iteración con bases de datos: {self.selected_databases}")  # Depuración

        # Preparar para nueva iteración
        self.current_database_index = 0
        self.database_decisions = []
        
        try:
            # Cargar el primer archivo
            current_database = self.selected_databases[self.current_database_index]
            database_name = current_database.split(" (")[0]
            database_path = current_database[current_database.find("(")+1:current_database.find(")")]
            
            # Cargar el archivo en la tercera página
            self.third_page.load_file(database_path)
            
            # Avanzar a la tercera página
            self.stack.setCurrentIndex(3)
        except Exception as e:
            print(f"Error en start_iteration: {e}")
            QMessageBox.warning(self, "Error", "Hubo un problema al iniciar el proceso. Por favor, intente nuevamente.")

    def go_to_final_page(self):
        """Ir a la página final para revisar decisiones."""
        try:
            self.final_page.load_decisions(self.selected_template, self.database_decisions)
            self.stack.setCurrentIndex(5)  # Cambiar al índice 5 (FinalPage)
        except Exception as e:
            print(f"Error al ir a la página final: {e}")
            QMessageBox.warning(self, "Error", "Hubo un problema al mostrar la página final. Por favor, intente nuevamente.")

if __name__ == "__main__":
    app = QApplication(sys.argv)
    window = MainWindow()
    window.show()
    sys.exit(app.exec_())